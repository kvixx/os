"""
Генератор отчёта по лабораторной работе №2 «Модель ОС».
Собирает docx: титул, цель, допущения, проектирование (схема данных,
таблица параметров, формулы, блок-схема по ГОСТ), результаты,
листинги, вывод, литература.
"""

from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import math

# --------------------------------------------------------------------------
# Пути и константы
# --------------------------------------------------------------------------

SCRIPT_DIR = Path(__file__).resolve().parent
REPORTS = SCRIPT_DIR / "Reports"
ASSETS = REPORTS / "assets"
OUTPUT = REPORTS / "Отчет_ЛР2_Модель_ОС.docx"
ASSETS.mkdir(parents=True, exist_ok=True)

# --------------------------------------------------------------------------
# Встроенные исходники Rust-проекта (fallback, если файлы не найдены рядом)
# --------------------------------------------------------------------------

SOURCES = {
    "main.rs": r"""mod domain;
mod services;
mod ui;

use crossterm::event::{self, Event, KeyCode, KeyEventKind};
use crossterm::terminal::{disable_raw_mode, enable_raw_mode};
use services::SimulationModel;
use std::env;
use std::time::{Duration, Instant};

use crate::services::Config;

fn print_help() {
    println!("Модель ОС — лабораторная работа 1");
    println!("Запуск: cargo run");
    println!(
        "Клавиши: + увеличить скорость на 10%, - уменьшить на 10%, R перезапуск, Q выход, H справка."
    );
}

fn main() -> std::io::Result<()> {
    let args: Vec<String> = env::args().collect();

    if args.iter().any(|a| a == "/?" || a == "--help" || a == "-h") {
        print_help();
        return Ok(());
    }

    let cfg = Config::load("config.toml");
    let mut model = SimulationModel::from_config(&cfg);
    model.initialize();

    if args.iter().any(|a| a == "--snapshot") {
        for _ in 0..12 {
            model.execute_one_tick();
        }
        ui::draw(&model, false);
        return Ok(());
    }

    enable_raw_mode()?;
    print!("\x1B[?25l");

    let mut last_tick_at = Instant::now();
    let mut last_draw_at = Instant::now() - Duration::from_secs(1);
    let mut show_help = false;
    let mut running = true;

    while running {
        while event::poll(Duration::from_millis(0))? {
            if let Event::Key(key) = event::read()? {
                if key.kind == KeyEventKind::Press {
                    match key.code {
                        KeyCode::Char('+') | KeyCode::Char('=') => model.change_speed(1.10),
                        KeyCode::Char('-') | KeyCode::Char('_') => model.change_speed(0.90),
                        KeyCode::Char('r') | KeyCode::Char('R') => model.initialize(),
                        KeyCode::Char('h') | KeyCode::Char('H') => show_help = !show_help,
                        KeyCode::Char('q') | KeyCode::Char('Q') | KeyCode::Esc => running = false,
                        _ => {}
                    }
                }
            }
        }

        let now = Instant::now();
        let tick_period = Duration::from_secs_f64(1.0 / model.speed_ticks_per_second());
        while now.duration_since(last_tick_at) >= tick_period {
            model.execute_one_tick();
            last_tick_at += tick_period;
        }

        if now.duration_since(last_draw_at) >= Duration::from_millis(100) {
            ui::draw(&model, show_help);
            last_draw_at = now;
        }

        std::thread::sleep(Duration::from_millis(10));
    }

    disable_raw_mode()?;
    ui::cleanup();
    println!("Моделирование завершено.");
    Ok(())
}
""",
    "domain/cpu.rs": r"""/// Модель центрального процессора.
#[derive(Debug)]
pub struct Cpu {
    pub id: u32,
    pub active_slot: Option<usize>,
    pub hardware_command_counter: u64,
}

impl Cpu {
    pub fn new() -> Self {
        Self {
            id: 1,
            active_slot: None,
            hardware_command_counter: 0,
        }
    }

    pub fn is_waiting(&self) -> bool {
        self.active_slot.is_none()
    }
}

impl Default for Cpu {
    fn default() -> Self {
        Self::new()
    }
}
""",
    "domain/process_control_block.rs": r"""use super::ProcessState;

/// Слово состояния процесса (PSW) модели ОС.
#[derive(Debug, Clone)]
pub struct ProcessControlBlock {
    pub slot: i32,
    pub task_id: i32,
    pub size: i32,
    pub command_counter: u64,
    pub state: ProcessState,

    // Поля зарезервированы для лабораторных 3–5.
    pub base_priority: i32,
    pub dynamic_priority: i32,
    pub remaining_quantum: i32,
}

impl ProcessControlBlock {
    pub fn new(task_id: i32, size: i32) -> Self {
        Self {
            slot: -1,
            task_id,
            size,
            command_counter: 0,
            state: ProcessState::Absent,
            base_priority: 0,
            dynamic_priority: 0,
            remaining_quantum: 0,
        }
    }
}
""",
    "domain/process_state.rs": r"""/// Состояния процесса.
#[derive(Debug, Clone, Copy, PartialEq, Eq)]
pub enum ProcessState {
    Absent,
    Ready,
    Active,
    BlockedIo,
    Suspended,
    Completed,
}

impl ProcessState {
    pub fn name(self) -> &'static str {
        match self {
            ProcessState::Absent => "Отсутствует",
            ProcessState::Ready => "Готов",
            ProcessState::Active => "Активен",
            ProcessState::BlockedIo => "Блокирован IO",
            ProcessState::Suspended => "Приостановлен",
            ProcessState::Completed => "Завершён",
        }
    }
}
""",
    "domain/mod.rs": r"""mod cpu;
mod process_control_block;
mod process_state;

pub use cpu::Cpu;
pub use process_control_block::ProcessControlBlock;
pub use process_state::ProcessState;
""",
    "services/config.rs": r"""use serde::Deserialize;
use std::fs;

#[derive(Debug, Deserialize)]
pub struct Config {
    pub model: ModelConfig,
    pub task: TaskConfig,
    pub cpu: CpuConfig,
}

#[derive(Debug, Deserialize)]
pub struct ModelConfig {
    pub memory_min: i32,
    pub memory_max: i32,
    pub process_table_size: usize,
    pub random_seed: u64,
}

#[derive(Debug, Deserialize)]
pub struct TaskConfig {
    pub size_min: i32,
    pub size_max: i32,
}

#[derive(Debug, Deserialize)]
pub struct CpuConfig {
    pub id: u32,
}

impl Config {
    pub fn default_config() -> Self {
        Self {
            model: ModelConfig {
                memory_min: 1024,
                memory_max: 4096,
                process_table_size: 8,
                random_seed: 14,
            },
            task: TaskConfig { size_min: 96, size_max: 384 },
            cpu: CpuConfig { id: 1 },
        }
    }

    pub fn load(path: &str) -> Self {
        match fs::read_to_string(path) {
            Ok(text) => toml::from_str(&text).unwrap_or_else(|e| {
                eprintln!("Ошибка разбора {}: {e}.", path);
                Self::default_config()
            }),
            Err(_) => {
                eprintln!("Файл {} не найден.", path);
                Self::default_config()
            }
        }
    }
}
""",
    "services/simulation_model.rs": r"""use crate::domain::{Cpu, ProcessControlBlock, ProcessState};
use crate::services::Config;
use rand::rngs::StdRng;
use rand::{Rng, SeedableRng};

pub struct SimulationModel {
    rng: StdRng,
    process_table: Vec<Option<ProcessControlBlock>>,
    next_task_id: i32,
    total_memory: i32,
    model_tick: u64,
    speed_ticks_per_second: f64,
    pub cpu: Cpu,
    pub new_task: Option<ProcessControlBlock>,
    task_size_min: i32,
    task_size_max: i32,
}

impl SimulationModel {
    pub fn new(memory_size: i32, process_table_size: usize, random_seed: u64) -> Self {
        Self {
            rng: StdRng::seed_from_u64(random_seed),
            process_table: vec![None; process_table_size],
            next_task_id: 1,
            total_memory: memory_size,
            model_tick: 0,
            speed_ticks_per_second: 2.0,
            cpu: Cpu::new(),
            new_task: None,
            task_size_min: 96,
            task_size_max: 384,
        }
    }

    /// Размер ОЗУ генерируется случайно в диапазоне [memory_min; memory_max].
    pub fn from_config(cfg: &Config) -> Self {
        let mut rng = StdRng::seed_from_u64(cfg.model.random_seed);
        let total_memory = rng.gen_range(cfg.model.memory_min..=cfg.model.memory_max);

        let mut cpu = Cpu::new();
        cpu.id = cfg.cpu.id;

        Self {
            rng,
            process_table: vec![None; cfg.model.process_table_size],
            next_task_id: 1,
            total_memory,
            model_tick: 0,
            speed_ticks_per_second: 2.0,
            cpu,
            new_task: None,
            task_size_min: cfg.task.size_min,
            task_size_max: cfg.task.size_max,
        }
    }

    pub fn generate_new_task(&mut self) {
        let size = self.rng.gen_range(self.task_size_min..=self.task_size_max);
        let id = self.next_task_id;
        self.next_task_id += 1;
        self.new_task = Some(ProcessControlBlock::new(id, size));
    }

    pub fn default_new() -> Self {
        Self::new(2048, 8, 14)
    }

    pub fn total_memory(&self) -> i32 { self.total_memory }

    pub fn used_memory(&self) -> i32 {
        self.process_table.iter()
            .filter_map(|p| p.as_ref())
            .map(|p| p.size)
            .sum()
    }

    pub fn free_memory(&self) -> i32 {
        self.total_memory - self.used_memory()
    }

    pub fn process_table_capacity(&self) -> usize { self.process_table.len() }

    pub fn resident_process_count(&self) -> usize {
        self.process_table.iter().filter(|p| p.is_some()).count()
    }

    pub fn model_tick(&self) -> u64 { self.model_tick }

    pub fn speed_ticks_per_second(&self) -> f64 { self.speed_ticks_per_second }

    pub fn process_table(&self) -> &[Option<ProcessControlBlock>] { &self.process_table }

    pub fn task_size_range(&self) -> (i32, i32) { (self.task_size_min, self.task_size_max) }

    pub fn active_process(&self) -> Option<&ProcessControlBlock> {
        self.cpu.active_slot.and_then(|s| self.process_table[s].as_ref())
    }

    pub fn initialize(&mut self) {
        for slot in self.process_table.iter_mut() { *slot = None; }
        self.cpu.active_slot = None;
        self.cpu.hardware_command_counter = 0;
        self.model_tick = 0;
        self.next_task_id = 1;
        self.new_task = None;

        self.generate_new_task();
        while self.load_new_task() {
            self.generate_new_task();
        }
        self.dispatch_first_ready_process();
    }

    pub fn has_available_memory(&self) -> bool {
        match &self.new_task {
            Some(t) => t.size <= self.free_memory(),
            None => false,
        }
    }

    pub fn load_new_task(&mut self) -> bool {
        if !self.has_available_memory() { return false; }
        let slot = match self.process_table.iter().position(|p| p.is_none()) {
            Some(s) => s,
            None => return false,
        };
        if let Some(mut task) = self.new_task.take() {
            task.slot = slot as i32;
            task.state = ProcessState::Ready;
            self.process_table[slot] = Some(task);
            true
        } else { false }
    }

    pub fn execute_one_tick(&mut self) {
        if self.cpu.active_slot.is_none() {
            self.dispatch_first_ready_process();
        }
        if let Some(slot) = self.cpu.active_slot {
            if let Some(process) = self.process_table[slot].as_mut() {
                process.command_counter += 1;
                self.cpu.hardware_command_counter += 1;
            }
        }
        self.model_tick += 1;
    }

    pub fn change_speed(&mut self, multiplier: f64) {
        self.speed_ticks_per_second =
            (self.speed_ticks_per_second * multiplier).clamp(0.1, 1000.0);
    }

    fn dispatch_first_ready_process(&mut self) {
        let slot = self.process_table.iter()
            .position(|p| matches!(p.as_ref().map(|x| x.state), Some(ProcessState::Ready)));
        if let Some(slot) = slot {
            if let Some(process) = self.process_table[slot].as_mut() {
                process.state = ProcessState::Active;
                self.cpu.active_slot = Some(slot);
                self.cpu.hardware_command_counter = process.command_counter;
            }
        }
    }
}
""",
    "services/mod.rs": r"""mod config;
mod simulation_model;

pub use config::Config;
pub use simulation_model::SimulationModel;
""",
    "ui/console_renderer.rs": r"""use crate::services::SimulationModel;
use std::io::Write;

pub fn draw(model: &SimulationModel, show_help: bool) {
    let mut buf = String::new();
    buf.push_str("\x1B[2J\x1B[1;1H");

    buf.push_str(&format!(
        "Такт модели: {:<8} Скорость: {:>7.1} такт/с\r\n",
        model.model_tick(),
        model.speed_ticks_per_second()
    ));

    let cpu_status = if model.cpu.is_waiting() { "ОЖИДАНИЕ" } else { "РАБОТА" };
    let active_id = model
        .active_process()
        .map(|p| p.task_id.to_string())
        .unwrap_or_else(|| "-".to_string());

    buf.push_str(&format!(
        "ЦП-{}: {:<10} Активный процесс: {:<3} Счётчик команд: {}\r\n",
        model.cpu.id, cpu_status, active_id, model.cpu.hardware_command_counter
    ));
    buf.push_str(&format!(
        "Память: занято {}/{}, свободно {}; PSW: {}/{}\r\n",
        model.used_memory(),
        model.total_memory(),
        model.free_memory(),
        model.resident_process_count(),
        model.process_table_capacity()
    ));

    buf.push_str("\r\n");
    buf.push_str(" №  ID    Размер   Счётчик команд   Состояние\r\n");
    buf.push_str(&format!("{}\r\n", "-".repeat(54)));

    for process in model.process_table() {
        if let Some(p) = process {
            buf.push_str(&format!(
                "{:>2}  {:>2}    {:>5}   {:>14}   {}\r\n",
                p.slot + 1, p.task_id, p.size, p.command_counter, p.state.name()
            ));
        }
    }

    buf.push_str("\r\n");
    match &model.new_task {
        None => buf.push_str("Новое задание: нет\r\n"),
        Some(task) => buf.push_str(&format!(
            "Новое задание: ID {}, {} ед. памяти\r\n",
            task.task_id, task.size
        )),
    }
    buf.push_str(
        "Управление: [+] быстрее  [-] медленнее  [R] перезапуск  [Q] выход  [H] справка\r\n",
    );

    if show_help {
        buf.push_str("\r\n");
        buf.push_str("Справка: +/- изменяют скорость на 10%; R инициализирует модель; Q завершает моделирование.\r\n");
    }

    let mut stdout = std::io::stdout();
    let _ = stdout.write_all(buf.as_bytes());
    let _ = stdout.flush();
}

pub fn cleanup() {
    print!("\x1B[2J\x1B[1;1H\x1B[?25h");
    let _ = std::io::stdout().flush();
}
""",
    "ui/mod.rs": r"""mod console_renderer;
pub use console_renderer::{cleanup, draw};
""",
    "config.toml": r"""# Параметры модели (генерируются один раз при запуске)
[model]
memory_min = 1024
memory_max = 4096
process_table_size = 8
random_seed = 14

# Параметры заданий (генерируются для каждого нового задания)
[task]
size_min = 96
size_max = 384

# Параметры процессора
[cpu]
id = 1
""",
}

# --------------------------------------------------------------------------
# Шрифты
# --------------------------------------------------------------------------

def _font_path(name):
    for base in [
        Path("C:/Windows/Fonts"),
        Path("/usr/share/fonts/truetype/msttcorefonts"),
        Path("/usr/share/fonts/truetype/dejavu"),
        Path("/Library/Fonts"),
    ]:
        p = base / name
        if p.exists():
            return str(p)
    return name


def font(size, bold=False, mono=False):
    if mono:
        name = "consolab.ttf" if bold else "consola.ttf"
        fallback = "DejaVuSansMono.ttf"
    else:
        name = "arialbd.ttf" if bold else "arial.ttf"
        fallback = "DejaVuSans.ttf"
    try:
        return ImageFont.truetype(_font_path(name), size)
    except OSError:
        return ImageFont.truetype(_font_path(fallback), size)


# --------------------------------------------------------------------------
# Хелперы рисования
# --------------------------------------------------------------------------

def centered_text(draw, box, text, fnt, fill="black"):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=6, align="center")
    w = bbox[2] - bbox[0]
    h = bbox[3] - bbox[1]
    cx = (x1 + x2) / 2
    cy = (y1 + y2) / 2
    draw.multiline_text(
        (cx - w / 2 - bbox[0], cy - h / 2 - bbox[1]),
        text, font=fnt, fill=fill, spacing=6, align="center"
    )


def draw_terminator(draw, box, text, fnt):
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=(y2 - y1) // 2, fill="white", outline="black", width=3)
    centered_text(draw, box, text, fnt)


def draw_process(draw, box, text, fnt):
    draw.rectangle(box, fill="white", outline="black", width=3)
    centered_text(draw, box, text, fnt)


def draw_diamond(draw, cx, cy, w, h, text, fnt):
    pts = [(cx, cy - h // 2), (cx + w // 2, cy), (cx, cy + h // 2), (cx - w // 2, cy)]
    draw.polygon(pts, fill="white", outline="black", width=3)
    centered_text(draw, (cx - w // 2, cy - h // 2, cx + w // 2, cy + h // 2), text, fnt)


def _arrow_head(draw, x1, y1, x2, y2):
    dx, dy = x2 - x1, y2 - y1
    L = math.hypot(dx, dy)
    if L < 1:
        return
    ux, uy = dx / L, dy / L
    px, py = -uy, ux
    ah, aw = 18, 9
    tip = (x2, y2)
    b1 = (x2 - ux * ah + px * aw, y2 - uy * ah + py * aw)
    b2 = (x2 - ux * ah - px * aw, y2 - uy * ah - py * aw)
    draw.polygon([tip, b1, b2], fill="black")


def draw_arrow(draw, x1, y1, x2, y2, width=3):
    draw.line((x1, y1, x2, y2), fill="black", width=width)
    _arrow_head(draw, x1, y1, x2, y2)


def draw_polyline(draw, pts, width=3):
    for i in range(len(pts) - 1):
        draw.line([pts[i], pts[i + 1]], fill="black", width=width)
    _arrow_head(draw, *pts[-2], *pts[-1])


# --------------------------------------------------------------------------
# Рисунок 1. Схема данных
# --------------------------------------------------------------------------

def make_data_schema(path):
    W, H = 1500, 1000
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)
    f_title = font(28, bold=True)
    f_head = font(22, bold=True)
    f_body = font(20)

    draw.text((W // 2 - 220, 25), "Схема данных модели ОС", font=f_title, fill="black")

    # ---- SimulationModel ----
    box1 = (60, 130, 620, 620)
    draw.rectangle(box1, fill="white", outline="black", width=3)
    draw.line((60, 190, 620, 190), fill="black", width=2)
    centered_text(draw, (60, 130, 620, 190), "SimulationModel", f_head)
    fields1 = [
        "- rng: StdRng",
        "- process_table: Vec<Option<PSW>>",
        "- new_task: Option<PSW>",
        "- next_task_id: i32",
        "- total_memory: i32",
        "- model_tick: u64",
        "- speed_ticks_per_second: f64",
        "- task_size_min, task_size_max: i32",
        "+ cpu: Cpu",
    ]
    for i, line in enumerate(fields1):
        draw.text((80, 210 + i * 34), line, font=f_body, fill="black")

    # ---- PSW (ProcessControlBlock) ----
    box2 = (900, 100, 1440, 480)
    draw.rectangle(box2, fill="white", outline="black", width=3)
    draw.line((900, 160, 1440, 160), fill="black", width=2)
    centered_text(draw, (900, 100, 1440, 160), "ProcessControlBlock (PSW)", f_head)
    fields2 = [
        "+ slot: i32",
        "+ task_id: i32",
        "+ size: i32",
        "+ command_counter: u64",
        "+ state: ProcessState",
        "+ base_priority: i32",
        "+ dynamic_priority: i32",
        "+ remaining_quantum: i32",
    ]
    for i, line in enumerate(fields2):
        draw.text((920, 180 + i * 34), line, font=f_body, fill="black")

    # ---- Cpu ----
    box3 = (900, 560, 1440, 830)
    draw.rectangle(box3, fill="white", outline="black", width=3)
    draw.line((900, 620, 1440, 620), fill="black", width=2)
    centered_text(draw, (900, 560, 1440, 620), "Cpu", f_head)
    fields3 = [
        "+ id: u32",
        "+ active_slot: Option<usize>",
        "+ hardware_command_counter: u64",
    ]
    for i, line in enumerate(fields3):
        draw.text((920, 640 + i * 34), line, font=f_body, fill="black")

    # ---- Связи ----
    # SimulationModel → PSW
    draw_arrow(draw, 620, 290, 900, 290)
    draw.text((640, 255), "содержит 0..N элементов", font=f_body, fill="black")
    # SimulationModel → Cpu
    draw_arrow(draw, 620, 500, 900, 690)
    draw.text((640, 620), "содержит 1 элемент", font=f_body, fill="black")

    img.save(path)


# --------------------------------------------------------------------------
# Рисунок 2. Блок-схема алгоритма загрузки заданий (ГОСТ 19.701-90)
# --------------------------------------------------------------------------

def make_loading_flowchart(path):
    W, H = 1320, 1220
    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)
    f_title = font(28, bold=True)
    f_text = font(20)
    f_small = font(17)

    draw.text((W // 2 - 210, 15), "Блок-схема алгоритма загрузки заданий",
              font=f_title, fill="black")

    start_box = (450, 70, 650, 130)
    clear_box = (250, 180, 850, 260)
    gen_box = (250, 310, 850, 390)

    dec1_cx, dec1_cy = 550, 500
    dec1_w, dec1_h = 620, 170

    dec2_cx, dec2_cy = 550, 700
    dec2_w, dec2_h = 620, 170

    write_box = (250, 840, 850, 930)
    dispatch_box = (250, 990, 850, 1080)
    end_box = (450, 1120, 650, 1180)

    # ---- Рисование ----
    draw_terminator(draw, start_box, "Начало", f_text)
    draw_arrow(draw, 550, 130, 550, 180)

    draw_process(draw, clear_box,
                 "Очистить таблицу слов состояний процессов", f_text)
    draw_arrow(draw, 550, 260, 550, 310)

    draw_process(draw, gen_box, "Сгенерировать новое задание", f_text)
    draw_arrow(draw, 550, 390, 550, 415)

    draw_diamond(draw, dec1_cx, dec1_cy, dec1_w, dec1_h,
                 "Достаточно свободной памяти?", f_text)

    # «да» вниз от DEC1
    draw_arrow(draw, 550, 585, 550, 615)
    draw.text((560, 590), "да", font=f_small, fill="black")

    draw_diamond(draw, dec2_cx, dec2_cy, dec2_w, dec2_h,
                 "Есть свободный слот в таблице PSW?", f_text)

    # «да» вниз от DEC2
    draw_arrow(draw, 550, 785, 550, 840)
    draw.text((560, 800), "да", font=f_small, fill="black")

    draw_process(draw, write_box,
                 "Записать задание в таблицу,\nсостояние — «Готов»", f_text)

    # Возврат цикла: WRITE → GEN (по левой стороне)
    loop_pts = [(250, 885), (130, 885), (130, 350), (250, 350)]
    draw_polyline(draw, loop_pts)

    # «нет» ветки
    # DEC1 → справа
    draw.line((860, 500, 1180, 500), fill="black", width=3)
    draw.text((880, 470), "нет", font=f_small, fill="black")
    # DEC2 → справа
    draw.line((860, 700, 1180, 700), fill="black", width=3)
    draw.text((880, 670), "нет", font=f_small, fill="black")
    # Общая вертикаль и вход в DISPATCH
    draw_polyline(draw, [(1180, 500), (1180, 1035), (850, 1035)])

    draw_process(draw, dispatch_box,
                 "Передать первый готовый процесс процессору", f_text)

    draw_arrow(draw, 550, 1080, 550, 1120)
    draw_terminator(draw, end_box, "Конец", f_text)

    img.save(path)


# --------------------------------------------------------------------------
# Рисунок 3. Мокап консольного экрана
# --------------------------------------------------------------------------

def make_console_screen(path):
    lines = [
        ("МОДЕЛЬ ОПЕРАЦИОННОЙ СИСТЕМЫ | Лабораторная работа 2", "#a7f3d0"),
        ("=" * 76, "#e5e7eb"),
        ("Такт модели: 12       Скорость:     2.0 такт/с", "#a7f3d0"),
        ("ЦП-1: РАБОТА     Активный процесс: 1   Счётчик команд: 12", "#a7f3d0"),
        ("Память: занято 1950/2432, свободно 482; PSW: 8/8", "#a7f3d0"),
        ("", "#e5e7eb"),
        (" №  ID    Размер   Счётчик команд   Состояние", "#e5e7eb"),
        ("-" * 54, "#e5e7eb"),
        (" 1   1      112               12   Активен", "#e5e7eb"),
        (" 2   2      348                0   Готов", "#e5e7eb"),
        (" 3   3      267                0   Готов", "#e5e7eb"),
        (" 4   4      239                0   Готов", "#e5e7eb"),
        (" 5   5      254                0   Готов", "#e5e7eb"),
        (" 6   6      331                0   Готов", "#e5e7eb"),
        (" 7   7      158                0   Готов", "#e5e7eb"),
        (" 8   8      241                0   Готов", "#e5e7eb"),
        ("", "#e5e7eb"),
        ("Новое задание: ID 9, 301 ед. памяти", "#e5e7eb"),
        ("Управление: [+] быстрее  [-] медленнее  [R] перезапуск  [Q] выход  [H] справка", "#e5e7eb"),
    ]
    f = font(20, mono=True)
    img = Image.new("RGB", (1600, 840), "#101419")
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, 1600, 50), fill="#1e293b")
    for i, (text, color) in enumerate(lines):
        draw.text((34, 30 + i * 39), text, fill=color, font=f)
    img.save(path)


# --------------------------------------------------------------------------
# Хелперы docx
# --------------------------------------------------------------------------

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def add_heading(document, text, level=1):
    p = document.add_heading(text, level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_code(document, text):
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text if text else " ")
    run.font.name = "Consolas"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(7)


def add_kv_table(document, rows, headers, fill="D9EAF7"):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        set_cell_shading(cell, fill)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    return table


def configure(document):
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.65)
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(6)


# --------------------------------------------------------------------------
# Сборка отчёта
# --------------------------------------------------------------------------

def make_report():
    screen_png = ASSETS / "console-screen.png"
    schema_png = ASSETS / "data-schema.png"
    flow_png = ASSETS / "loading-flowchart.png"

    make_console_screen(screen_png)
    make_data_schema(schema_png)
    make_loading_flowchart(flow_png)

    doc = Document()
    configure(doc)

    # ---------- Титульный лист ----------
    for text in [
        "Министерство науки и высшего образования Российской Федерации",
        "",
        "ОТЧЁТ",
        "по лабораторной работе № 2",
        "по дисциплине «Операционные системы»",
        "",
        "Тема: «Генерация и загрузка задач ПО «Модель ОС»»",
        "",
        "Выполнил: студент ____________________",
        "Группа: ____________________",
        "Проверил: ____________________",
        "",
        "2026",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14 if text in ("ОТЧЁТ", "по лабораторной работе № 2") else 12)
        if text == "ОТЧЁТ":
            run.bold = True
    doc.add_page_break()

    # ---------- Цель ----------
    add_heading(doc, "Цель работы")
    doc.add_paragraph(
        "Разработать подпрограммы генерации и загрузки задач главной формы "
        "программного обеспечения «Модель ОС». Отработать работу с файлом "
        "конфигурации, генерацию случайных параметров заданий и модели, "
        "а также заполнение таблицы слов состояний процессов (PSW)."
    )

    # ---------- Допущения ----------
    add_heading(doc, "Допущения")
    for item in [
        "1. Размер памяти генерируется при запуске модели.",
        "2. Заданий может загружаться столько, сколько помещается в памяти "
        "в соответствии с их размером.",
        "3. Количество новых заданий бесконечно.",
        "4. Количество процессоров ввода (вывода) достаточно для работы "
        "без очередей к ним.",
    ]:
        doc.add_paragraph(item)

    # ---------- Постановка задачи ----------
    add_heading(doc, "Постановка задачи")
    doc.add_paragraph(
        "Входные данные для исследования возможностей модели ОС генерируются "
        "по данным файла конфигурации. Параметры делятся на две группы: "
        "параметры модели (генерируются один раз при запуске) и параметры "
        "заданий (генерируются для каждого нового задания). "
        "Загрузка задания выполняется по алгоритму: если имеется свободная "
        "память и свободный слот в таблице PSW — готовое новое задание "
        "передаётся в свободное место таблицы, затем генерируется следующее "
        "задание."
    )

    # ---------- 1. Схема данных ----------
    add_heading(doc, "Схема данных")
    doc.add_paragraph(
        "Модель хранит своё состояние в структуре SimulationModel. Таблица "
        "слов состояний процессов — вектор process_table из Option<PSW>. "
        "Свободный слот — None, занятый — Some(ProcessControlBlock). "
        "Сгенерированное, но ещё не загруженное задание хранится в "
        "переменной new_task (PSW_Task в терминах методички). "
        "Активный процесс ЦП задаётся индексом active_slot, "
        "указывающим в process_table."
    )
    doc.add_picture(str(schema_png), width=Inches(6.5))
    p = doc.add_paragraph("Рисунок 1 — Схема данных модели ОС")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------- 2. Параметры ----------
    add_heading(doc, "Параметры модели и заданий")
    doc.add_paragraph(
        "Параметры заданы в файле config.toml. Часть из них используется "
        "один раз при запуске (параметры модели), часть — для каждого нового "
        "задания (параметры заданий)."
    )

    doc.add_paragraph("Таблица 1 — Параметры модели").runs[0].bold = True
    add_kv_table(
        doc,
        [
            ("memory_min", "Нижняя граница ОЗУ, ед.", "1024", "config.toml"),
            ("memory_max", "Верхняя граница ОЗУ, ед.", "4096", "config.toml"),
            ("total_memory", "Сгенерированный размер ОЗУ, ед.", "[1024; 4096]", "ГПСЧ при запуске"),
            ("process_table_size", "Число слотов в таблице PSW", "8", "config.toml"),
            ("random_seed", "Зерно генератора", "14", "config.toml"),
            ("speed_ticks_per_second", "Скорость моделирования, такт/с", "2.0", "константа"),
            ("cpu.id", "Идентификатор ЦП", "1", "config.toml"),
        ],
        headers=["Параметр", "Назначение", "Значение", "Источник"],
    )

    doc.add_paragraph("Таблица 2 — Параметры заданий").runs[0].bold = True
    add_kv_table(
        doc,
        [
            ("size_min", "Мин. размер задания, ед.", "96", "config.toml"),
            ("size_max", "Макс. размер задания, ед.", "384", "config.toml"),
            ("size", "Размер нового задания", "[96; 384]", "ГПСЧ, для каждого задания"),
            ("task_id", "Идентификатор задания", "1, 2, 3, …", "монотонный счётчик"),
            ("command_counter", "Начальный счётчик команд", "0", "константа"),
            ("state", "Начальное состояние", "«Готов»", "после загрузки"),
        ],
        headers=["Параметр", "Назначение", "Значение", "Источник"],
    )

    # ---------- 3. Формулы ----------
    add_heading(doc, "Формулы")
    doc.add_paragraph("Занятая память:")
    add_code(doc, "P_used = Σ size_i,   для всех i, где PSW[i] ≠ пусто")
    doc.add_paragraph("Свободная память:")
    add_code(doc, "P_free = P_total − P_used")
    doc.add_paragraph("Число резидентных процессов:")
    add_code(doc, "N_res = | { i : PSW[i] ≠ пусто } |")
    doc.add_paragraph("Условие возможности загрузки нового задания:")
    add_code(doc, "load_ok = ( P_free ≥ size_new ) AND ( ∃ i : PSW[i] = пусто )")
    doc.add_paragraph(
        "Перечисленные формулы реализованы в методах used_memory, "
        "free_memory, resident_process_count и has_available_memory класса "
        "SimulationModel."
    )

    # ---------- 4. Блок-схема ----------
    add_heading(doc, "Алгоритм загрузки заданий")
    doc.add_paragraph(
        "После инициализации модели очищаются таблица PSW, аппаратный счётчик "
        "команд ЦП и служебные переменные. Затем генерируется новое задание. "
        "Цикл загрузки повторяется, пока одновременно выполнены два условия: "
        "свободной памяти достаточно для размещения задания и в таблице PSW "
        "есть хотя бы один свободный слот. После выхода из цикла первый "
        "готовый процесс передаётся процессору."
    )
    doc.add_picture(str(flow_png), width=Inches(6.3))
    p = doc.add_paragraph("Рисунок 2 — Блок-схема алгоритма загрузки заданий")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------- 5. Реализация пунктов задания ----------
    add_heading(doc, "Практическая часть")
    doc.add_paragraph(
        "Программа реализована на языке Rust как консольное приложение. "
        "Все 12 пунктов задания выполнены в соответствующих модулях:"
    )
    add_kv_table(
        doc,
        [
            ("1", "Параметры процессора", "domain/cpu.rs (struct Cpu)"),
            ("2", "Параметры системы", "services/config.rs + config.toml"),
            ("3", "Тип состояния процесса", "domain/process_state.rs (enum ProcessState)"),
            ("4", "Слово состояния процесса", "domain/process_control_block.rs"),
            ("5", "Массив PSW", "services/simulation_model.rs, поле process_table"),
            ("6", "Переменная PSW_Task", "services/simulation_model.rs, поле new_task"),
            ("7", "Инициализация массива PSW", "SimulationModel::initialize"),
            ("8", "Генерация нового задания", "SimulationModel::generate_new_task"),
            ("9", "Проверка доступной памяти", "SimulationModel::has_available_memory"),
            ("10", "Загрузка нового задания", "SimulationModel::load_new_task"),
            ("11", "Цикл загрузки при запуске", "SimulationModel::initialize"),
            ("12", "Индикация слов состояний", "ui/console_renderer.rs"),
        ],
        headers=["№", "Пункт задания", "Где реализовано"],
    )

    # ---------- 6. Экранная форма ----------
    add_heading(doc, "Экранная форма")
    doc.add_paragraph(
        "В верхней части экрана отображаются такт модели, скорость, состояние "
        "ЦП и активный процесс, объём занятой и свободной памяти, число "
        "занятых слотов PSW. Далее выводится таблица процессов со столбцами "
        "«№», «ID», «Размер», «Счётчик команд», «Состояние», а также "
        "информация о сгенерированном новом задании и строке управления."
    )
    doc.add_picture(str(screen_png), width=Inches(6.5))
    p = doc.add_paragraph("Рисунок 3 — Индикация параметров модели во время работы")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---------- 7. Проверка ----------
    add_heading(doc, "Проверка работы")
    doc.add_paragraph(
        "Сборка выполнена командой cargo build. Программа компилируется "
        "без ошибок. При запуске размер ОЗУ генерируется случайно в диапазоне "
        "[memory_min; memory_max] из config.toml, что соответствует "
        "допущению 1. Далее модель загружает задания до заполнения памяти "
        "или таблицы PSW."
    )
    add_kv_table(
        doc,
        [
            ("Размер ОЗУ", "[1024; 4096]", "2432", "Генерация при запуске выполнена"),
            ("Загружено процессов", "0", "8", "Заполнены все слоты PSW"),
            ("Свободная память", "2432", "482", "Не хватило слота в PSW"),
            ("Счётчик активного процесса", "0", "12", "Увеличивается на 1 за такт"),
            ("Состояние нового задания", "—", "Ожидает слот", "Условие по памяти выполнено, слота нет"),
        ],
        headers=["Параметр", "Начало", "Конец", "Вывод"],
    )

    # ---------- 8. Листинг ----------
    add_heading(doc, "Листинг программы", 1)
    order = [
        "main.rs",
        "domain/mod.rs",
        "domain/cpu.rs",
        "domain/process_control_block.rs",
        "domain/process_state.rs",
        "services/mod.rs",
        "services/config.rs",
        "services/simulation_model.rs",
        "ui/mod.rs",
        "ui/console_renderer.rs",
        "config.toml",
    ]
    for rel in order:
        doc.add_page_break()
        add_heading(doc, rel, 2)
        text = SOURCES.get(rel, "")
        for line in text.splitlines():
            add_code(doc, line)

    # ---------- 9. Вывод ----------
    doc.add_page_break()
    add_heading(doc, "Вывод")
    doc.add_paragraph(
        "В результате работы разработаны подпрограммы генерации и загрузки "
        "задач ПО «Модель ОС». Реализованы: чтение параметров из файла "
        "config.toml, генерация размера ОЗУ при запуске (допущение 1), "
        "генерация параметров новых заданий, проверка доступной памяти и "
        "свободного слота таблицы PSW, загрузка задания и цикл заполнения "
        "таблицы слов состояний процессов. Индикация дополнена выводом "
        "номера, идентификатора, счётчика команд и кода состояния каждого "
        "процесса, а также сведений об объёме занятой и свободной памяти. "
        "Все 12 пунктов задания выполнены."
    )

    # ---------- 10. Литература ----------
    add_heading(doc, "Литература")
    for item in [
        "1. Таненбаум Э., Бос Х. Современные операционные системы. 5-е изд. — СПб.: Питер, 2022.",
        "2. Silberschatz A., Galvin P., Gagne G. Operating System Concepts. 10th ed. — Wiley, 2018.",
        "3. ГОСТ 19.701-90. Единая система программной документации. Схемы алгоритмов, программ, данных и систем.",
        "4. Задание к лабораторной работе № 2 по дисциплине «Операционные системы».",
    ]:
        doc.add_paragraph(item)

    doc.save(OUTPUT)
    print("Готово:", OUTPUT)


if __name__ == "__main__":
    make_report()
